from copy import copy
from pathlib import Path

import xlrd
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
SOURCE_PATH = BASE_DIR / "test.xls"
TEMPLATE_PATH = BASE_DIR / "模型分项性能与精度对比.xlsx"
OUTPUT_PATH = BASE_DIR / "模型分项性能与精度对比_已补充Qwen3-VL-8B.xlsx"


def text(sheet, row, column):
    return str(sheet.cell_value(row - 1, column - 1)).strip()


def number(sheet, row, column):
    return float(text(sheet, row, column))


source_book = xlrd.open_workbook(SOURCE_PATH)
source_sheet = source_book.sheet_by_name("优化后")

versions = [
    {
        "name": text(source_sheet, 1, 12),
        "columns": (12, 13, 14),
    },
    {
        "name": text(source_sheet, 1, 7),
        "columns": (6, 7, 8),
    },
    {
        "name": f"{text(source_sheet, 1, 2)} + {text(source_sheet, 1, 3)}",
        "columns": (2, 3, 4),
    },
]

baseline_ttft = {
    int(number(source_sheet, source_row, 1)): number(source_sheet, source_row, 12)
    for source_row in range(3, 7)
}

records = []
for version in versions:
    ttft_column, tps_column, decode_tps_column = version["columns"]
    for source_row in range(3, 7):
        test_index = int(number(source_sheet, source_row, 1))
        ttft = number(source_sheet, source_row, ttft_column)
        baseline = baseline_ttft[test_index]
        records.append(
            {
                "model": "Qwen3-VL-8B",
                "version": version["name"],
                "ttft": ttft,
                "speedup": baseline / ttft,
                "reduction": 1 - ttft / baseline,
                "tps": number(source_sheet, source_row, tps_column),
                "decode_tps": number(source_sheet, source_row, decode_tps_column),
                "test_index": test_index,
            }
        )

workbook = load_workbook(TEMPLATE_PATH)
style_sheet = workbook["ResNet50"]
if "Qwen3-VL-8B" in workbook.sheetnames:
    del workbook["Qwen3-VL-8B"]

target_sheet = workbook.copy_worksheet(style_sheet)
target_sheet.title = "Qwen3-VL-8B"

for merged_range in list(target_sheet.merged_cells.ranges):
    target_sheet.unmerge_cells(str(merged_range))
target_sheet.delete_rows(2, target_sheet.max_row - 1)
for row_index in list(target_sheet.row_dimensions):
    if row_index > 1:
        del target_sheet.row_dimensions[row_index]

headers = [
    "模型",
    "版本",
    "TTFT延迟 (s)",
    "TTFT加速比",
    "TTFT降低比",
    "吞吐指标",
    "吞吐数值 (toks/s)",
    "测试条件",
]
for column, value in enumerate(headers, start=1):
    target_sheet.cell(row=1, column=column, value=value)

for target_row, record in enumerate(records, start=2):
    for column in range(1, 9):
        source_cell = style_sheet.cell(row=2, column=column)
        target_cell = target_sheet.cell(row=target_row, column=column)
        target_cell._style = copy(source_cell._style)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = copy(source_cell.number_format)

    values = [
        record["model"],
        record["version"],
        record["ttft"],
        f'{record["speedup"]:.3f}×',
        f'{record["reduction"]:.2%}',
        "TPS / Decode TPS",
        f'{record["tps"]:.3f} / {record["decode_tps"]:.3f}',
        f'测试序号 {record["test_index"]}',
    ]
    for column, value in enumerate(values, start=1):
        target_sheet.cell(row=target_row, column=column, value=value)
    target_sheet.cell(row=target_row, column=3).number_format = "0.000"

footer_row = len(records) + 2
for column in range(1, 9):
    source_cell = style_sheet.cell(row=6, column=column)
    target_cell = target_sheet.cell(row=footer_row, column=column)
    target_cell._style = copy(source_cell._style)
    target_cell.alignment = copy(source_cell.alignment)
target_sheet.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=8)
target_sheet.cell(
    row=footer_row,
    column=1,
    value=(
        "数据来源: test.xls / 优化后；测试序号 0-3。加速比与延迟降低比按相同测试序号的"
        "未优化 TTFT 计算；源数据未提供精度指标。"
    ),
)
target_sheet.row_dimensions[footer_row].height = 32
target_sheet.print_area = f"A1:H{footer_row}"
target_sheet.freeze_panes = "A2"

workbook.active = workbook.sheetnames.index(target_sheet.title)
workbook.save(OUTPUT_PATH)
print(OUTPUT_PATH)
